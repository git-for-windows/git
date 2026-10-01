#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Fuellt die Word-Vorlage "BEI_BW B - Gesundheitsbogen" aus einer JSON-Zuordnung.

Aufruf:
    python fill_bei_bw_b.py VORLAGE.docx ZUORDNUNG.json AUSGABE.docx

Schema der JSON-Datei:
{
  "person": "Anna Apfelkuchen",
  "datum_gespraech": "XX.XX.XXXX",
  "quellen": {"aerztliche_unterlagen": null, "teilhabebericht": null, "vorheriger_gesundheitsbogen": null},
  "diagnosen": ["Entwicklungsstoerung ... (F70.0)"],
  "koerperfunktionen": {
    "b130": {"angekreuzt": true, "quelle": "Teilhabebericht 13.02.26", "erlaeuterungen": ["FS: ..."]}
  },
  "kapitel_zusaetzlich": [],
  "ergaenzende_hinweise": {"zitate": ["FS: ..."], "entwurf": "Entwurf, bitte pruefen: ..."},
  "nicht_zugeordnet": []
}

Das Skript setzt die Kaestchen der Kapitel-Uebersicht und der ICF-Code-Zeilen,
traegt Diagnosen als Listeneintraege ein, schreibt Erlaeuterungen in die
dritte Spalte der Code-Zeile und haengt die Ergaenzenden Hinweise an. Das
Layout der Vorlage bleibt unveraendert. Benoetigt: python-docx.
"""
import copy
import json
import re
import sys

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
except ImportError:  # pragma: no cover
    sys.stderr.write("Fehlt: python-docx. Installieren mit: pip install python-docx\n")
    sys.exit(2)

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
PLACEHOLDER = "XX.XX.XXXX"
KEINE_DIAGNOSE = "Aus den vorliegenden Unterlagen nicht ermittelbar – bitte aus der Akte ergänzen."


def q14(tag):
    return "{%s}%s" % (W14, tag)


# ---------------------------------------------------------------- Hilfsfunktionen
def ptext(p):
    return "".join(r.text for r in p.runs)


def is_heading(p):
    return (p.style is not None and p.style.name.startswith("BEI_BW")) or \
           (p._p.pPr is not None and p._p.pPr.pStyle is not None and
            str(p._p.pPr.pStyle.get(qn("w:val"))).startswith("BEIBW"))


def clean(items):
    out, seen = [], set()
    for it in items or []:
        s = (it or "").strip()
        if s and s not in seen:
            out.append(s)
            seen.add(s)
    return out


def set_text(p, text):
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        run = p.add_run(text)
        _copy_mark_rpr(p, run)


def _copy_mark_rpr(src_p, run):
    """Uebernimmt die Zeichenformatierung der Absatzmarke, wenn der Absatz keine Runs hat."""
    pPr = src_p._p.pPr
    if pPr is not None and pPr.find(qn("w:rPr")) is not None and run._r.rPr is None:
        rpr = copy.deepcopy(pPr.find(qn("w:rPr")))
        for bad in rpr.findall(qn("w:rStyle")):
            pass
        run._r.insert(0, rpr)


def clone_after(anchor, text, template=None):
    src = template if template is not None else anchor
    new_p = copy.deepcopy(src._p)
    for child in list(new_p):
        if child.tag != qn("w:pPr"):
            new_p.remove(child)
    anchor._p.addnext(new_p)
    np = Paragraph(new_p, anchor._parent)
    run = np.add_run(text)
    if src.runs and src.runs[0]._r.rPr is not None:
        run._r.insert(0, copy.deepcopy(src.runs[0]._r.rPr))
    else:
        _copy_mark_rpr(src, run)
    return np


def checkbox_sdts(element):
    return [s for s in element.iter(qn("w:sdt")) if s.find(".//" + q14("checkbox")) is not None]


def set_checkbox(sdt, checked):
    cb = sdt.find(".//" + q14("checkbox"))
    chk = cb.find(q14("checked"))
    if chk is None:
        chk = cb.makeelement(q14("checked"), {})
        cb.insert(0, chk)
    chk.set(q14("val"), "1" if checked else "0")
    state = cb.find(q14("checkedState") if checked else q14("uncheckedState"))
    code = state.get(q14("val")) if state is not None else ("2612" if checked else "2610")
    glyph = chr(int(code, 16))
    ts = sdt.findall(".//" + qn("w:t"))
    if ts:
        ts[0].text = glyph
        for t in ts[1:]:
            t.text = ""


def fill_cell(cell, items):
    """Schreibt Zeilen in eine Tabellenzelle; erster Absatz der Zelle ist die Formatvorlage."""
    items = clean(items)
    ps = cell.paragraphs
    base = ps[0]
    for p in ps[1:]:
        p._p.getparent().remove(p._p)
    if not items:
        set_text(base, "")
        return
    set_text(base, items[0])
    last = base
    for it in items[1:]:
        last = clone_after(last, it, base)


# ---------------------------------------------------------------- Bloecke
def fill_diagnosen(doc, data, log):
    diags = clean(data.get("diagnosen")) or [KEINE_DIAGNOSE]
    paras = doc.paragraphs
    idx = next((i for i, p in enumerate(paras) if ptext(p).strip().startswith("Diagnosen, die gesundheitliche")), None)
    if idx is None:
        raise SystemExit("Vorlage unerwartet: Einleitungssatz des Diagnoseblocks fehlt.")
    items = []
    j = idx + 1
    while j < len(paras) and not is_heading(paras[j]):
        items.append(paras[j])
        j += 1
    list_items = [p for p in items if p._p.pPr is not None and p._p.pPr.numPr is not None]
    if list_items:
        tpl = list_items[0]
        for p in list_items[1:]:
            p._p.getparent().remove(p._p)
        set_text(tpl, diags[0])
        last = tpl
        for d in diags[1:]:
            last = clone_after(last, d, tpl)
    else:
        # keine Listenabsaetze in der Vorlage: Eintraege als normale Absaetze nach der Einleitung
        last = paras[idx]
        for d in diags:
            last = clone_after(last, d, paras[idx])
    log.append("Diagnosen eingetragen: %d" % len(diags))


def fill_uebersicht(doc, chapters, log):
    found = set()
    for p in doc.paragraphs:
        sdts = checkbox_sdts(p._p)
        if not sdts:
            continue
        m = re.match(r"\s*([1-8])\s", ptext(p))
        if not m:
            continue
        k = int(m.group(1))
        found.add(k)
        set_checkbox(sdts[0], k in chapters)
    missing = sorted(chapters - found)
    if missing:
        log.append("Kapitel ohne Kaestchen in der Vorlage: %s" % ", ".join(map(str, missing)))
    log.append("Kapitel angekreuzt: %s" % (", ".join(map(str, sorted(chapters))) or "keine"))


def fill_codes(doc, data, log):
    kf = data.get("koerperfunktionen", {}) or {}
    wanted = {k.strip().lower(): v for k, v in kf.items()}
    seen = set()
    checked_chapters = set()
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            if len(cells) < 2:
                continue
            m = re.match(r"\s*(b\d{3})\b", cells[1].text)
            if not m:
                continue
            code = m.group(1).lower()
            spec = wanted.get(code) or {}
            sdts = checkbox_sdts(cells[0]._tc)
            if sdts:
                set_checkbox(sdts[0], bool(spec.get("angekreuzt")))
            if len(cells) >= 3 and cells[2]._tc is not cells[1]._tc:
                fill_cell(cells[2], spec.get("erlaeuterungen"))
            if spec.get("angekreuzt"):
                checked_chapters.add(int(code[1]))
            seen.add(code)
    missing = [c for c in wanted if c not in seen]
    if missing:
        log.append("Codes nicht in der Vorlage vorhanden (Tabelle des Kapitels fehlt): %s" % ", ".join(missing))
    derived = [c for c, s in wanted.items() if s.get("angekreuzt") and "abgeleitet" in str(s.get("quelle", "")).lower()]
    if derived:
        log.append("Kreuze aus Diagnosetext abgeleitet, nicht ausdruecklich belegt (bitte pruefen): %s" % ", ".join(derived))
    no_tick = [c for c, s in wanted.items() if c in seen and not s.get("angekreuzt") and clean(s.get("erlaeuterungen"))]
    if no_tick:
        log.append("Codes mit Hinweisen aus dem Gespraech, aber ohne Kreuz (fachliche Bewertung noetig): %s" % ", ".join(no_tick))
    return checked_chapters


def fill_hinweise(doc, data, datum, log):
    eh = data.get("ergaenzende_hinweise") or {}
    if isinstance(eh, list):
        eh = {"zitate": eh, "entwurf": ""}
    zitate = clean(eh.get("zitate"))
    entwurf = (eh.get("entwurf") or "").strip()
    lines = []
    if zitate:
        lines.append("Aus dem Gespräch am %s:" % datum)
        lines += zitate
    if entwurf:
        lines.append(entwurf)
    if not lines:
        return
    paras = doc.paragraphs
    heading = None
    for p in paras:
        if is_heading(p) and ptext(p).strip() == "Ergänzende Hinweise":
            heading = p
    if heading is None:
        raise SystemExit("Vorlage unerwartet: Ueberschrift 'Ergaenzende Hinweise' fehlt.")
    nxt = heading._p.getnext()
    tpl = None
    if nxt is not None and nxt.tag == qn("w:p"):
        cand = Paragraph(nxt, heading._parent)
        if not is_heading(cand):
            tpl = cand
    if tpl is None:
        tpl = next((p for p in paras if ptext(p).strip().startswith("Diagnosen, die gesundheitliche")), heading)
    last = heading
    # vorhandene leere Absaetze direkt nach der Ueberschrift als Platz nutzen
    slot = tpl if (tpl is not heading and not ptext(tpl).strip() and tpl._p.getprevious() is heading._p) else None
    for i, line in enumerate(lines):
        if i == 0 and slot is not None:
            set_text(slot, line)
            last = slot
        else:
            last = clone_after(last, line, tpl)
    log.append("Ergaenzende Hinweise: %d Zeilen" % len(lines))


def write_protocol(path, data, log):
    q = data.get("quellen") or {}
    lines = ["Pruefprotokoll Gesundheitsbogen (automatisch erzeugt)", ""]
    lines.append("Person: %s" % data.get("person", ""))
    lines.append("Gespraech am: %s" % (data.get("datum_gespraech") or PLACEHOLDER))
    lines.append("Aerztliche Unterlagen: %s" % (q.get("aerztliche_unterlagen") or "nicht beigefuegt"))
    lines.append("Teilhabebericht: %s" % (q.get("teilhabebericht") or "nicht beigefuegt"))
    lines.append("Vorheriger Gesundheitsbogen: %s" % (q.get("vorheriger_gesundheitsbogen") or "nicht beigefuegt"))
    lines.append("")
    lines += log
    lines.append("")
    lines.append("Nicht zugeordnete gesundheitsbezogene Saetze:")
    lines += ["- " + s for s in clean(data.get("nicht_zugeordnet"))] or ["- keine"]
    lines.append("")
    lines.append("Bitte pruefen: Diagnosen und Kreuze stammen nur aus medizinischen Unterlagen bzw. Teilhabebericht; aus dem Gespraech stammen nur Erlaeuterungen und Ergaenzende Hinweise.")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")


def main(argv):
    if len(argv) != 4:
        sys.stderr.write(__doc__)
        return 1
    template, mapping, output = argv[1], argv[2], argv[3]
    with open(mapping, encoding="utf-8") as fh:
        data = json.load(fh)
    datum = (data.get("datum_gespraech") or "").strip() or PLACEHOLDER
    log = []
    doc = Document(template)
    fill_diagnosen(doc, data, log)
    chapters = fill_codes(doc, data, log)
    chapters |= set(int(k) for k in (data.get("kapitel_zusaetzlich") or []) if str(k).isdigit())
    fill_uebersicht(doc, chapters, log)
    fill_hinweise(doc, data, datum, log)
    doc.save(output)
    proto = output.rsplit(".", 1)[0] + "_Pruefprotokoll.txt"
    write_protocol(proto, data, log)
    print("Geschrieben: %s" % output)
    print("Pruefprotokoll: %s" % proto)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
